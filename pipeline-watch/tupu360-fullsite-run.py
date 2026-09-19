#!/usr/bin/env python3
"""tupu360 (图谱天下) full-site collector -- resumable runner + table builder.

Scope: every tenant of the *public* ``careersite`` family that the 2026-09-19
census found to carry at least one real posting (60 companies x 3 channels =
180 units).  The ``wxtemp`` family (``<slug>.tupu360.com``) has no anonymous
public route and is not touched here.

Contract with the rest of the collector
---------------------------------------
One unit = one (tenant slug, scope).  For each unit this runner calls

    qiuzhao.collector.p1_platform_tupu360.collect(slug, scope, unit_dir)

and then the shared validator

    qiuzhao.collector.p1_pipeline.validate_result(payload, name, scope)

exactly like ``p1_pipeline.collect_process`` does for the daily run.  A unit is
finished only when the validated payload has been written to
``<out>/units/<slug>__<scope>/validated.json``; that file is also the resume
marker, so re-running the script after an interruption skips finished units.

Pacing (hard constraint, not configurable downwards)
----------------------------------------------------
``QIUZHAO_PLATFORM_REQUEST_INTERVAL`` defaults to 2.0 s here and every HTTP call
the adapter makes sleeps that long first, so per-worker spacing is >= 2.0 s.
``--workers`` defaults to 3 and is capped at 3, i.e. aggregate <= 1.5 req/s.

Honesty rules
-------------
Nothing is inferred.  A scope the official site answers with 0 postings is
recorded as 0; a unit that fails keeps its error string; ``partial`` stays
``partial``.  The table builder copies validated rows verbatim, it never fills a
missing deadline, cohort or headcount.

Usage
-----
    python3 pipeline-watch/tupu360-fullsite-run.py                 # collect + build tables
    python3 pipeline-watch/tupu360-fullsite-run.py --build-table   # rebuild tables from disk
    python3 pipeline-watch/tupu360-fullsite-run.py --only a,b --scopes social --force
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import threading
import time
import traceback
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

HERE = Path(__file__).resolve().parent
INPUTS = HERE / 'tupu360-fullsite' / 'inputs'
TABLE_DIR = HERE / 'tupu360-fullsite'
DEFAULT_OUT = Path('/Volumes/臭垃圾桶/生财MCP/_worktrees/tupu360-fullsite-out/tupu360-fullsite')
TABLE_STEM = 'tupu360-全站岗位-20260919'

SCOPES = ('campus', 'intern', 'social')
SCOPE_LABEL = {'campus': '校园招聘', 'intern': '实习招聘', 'social': '社会招聘'}
CENSUS_CHANNEL = {'campus': 'CAMPUSRECRUITMENT', 'intern': 'INTERNSHIPRECRUITMENT',
                  'social': 'SOCIALRECRUITMENT'}
MAX_WORKERS = 3
DEFAULT_INTERVAL = 2.0
EXCEL_CELL_LIMIT = 30000

# Extra targets the 2026-09-19 census could not classify.  They are re-checked by
# hand and their measured outcome is written into the receipt, so the document can
# state exactly what happened instead of "unknown".
RETRY_TARGETS = {
    'bbac': ('北京奔驰人力资源', 'https://careersite.tupu360.com/bbac/position/index'),
    'hellohr': ('三人行', 'https://www.hellohr.cn/'),
    'careerjnj': ('强生社招', 'https://career.jnj.com.cn/'),
}

_PRINT_LOCK = threading.Lock()


def log(message):
    stamp = datetime.now().strftime('%H:%M:%S')
    with _PRINT_LOCK:
        print(f'[{stamp}] {message}', flush=True)


# --------------------------------------------------------------------------- #
# inputs
# --------------------------------------------------------------------------- #
def load_inputs(inputs: Path):
    scope = json.loads((inputs / 'tupu_scope.json').read_text(encoding='utf-8'))
    census = json.loads((inputs / 'tupu_census.json').read_text(encoding='utf-8'))
    tenants = json.loads((inputs / 'tupu_tenants.json').read_text(encoding='utf-8'))
    return scope, census, tenants


def tenant_names(tenants):
    return {row[0].lstrip('/'): row[1] for row in tenants}


# --------------------------------------------------------------------------- #
# one unit
# --------------------------------------------------------------------------- #
def unit_dir(out: Path, slug: str, scope: str) -> Path:
    return out / 'units' / f'{slug}__{scope}'


def read_validated(path: Path):
    try:
        data = json.loads(path.read_text(encoding='utf-8'))
    except Exception:
        return None
    if not isinstance(data, dict) or 'validated' not in data:
        return None
    return data


def collect_unit(slug, name, scope, out: Path, force=False):
    """Collect + validate one (tenant, scope) unit.  Returns a ledger record."""
    from qiuzhao.collector import p1_pipeline as pipeline
    from qiuzhao.collector import p1_platform_tupu360 as tupu

    directory = unit_dir(out, slug, scope)
    directory.mkdir(parents=True, exist_ok=True)
    marker = directory / 'validated.json'
    cached = None if force else read_validated(marker)
    if cached is not None:
        record = dict(cached.get('meta') or {})
        record['resumed'] = True
        return record

    started = time.time()
    record = {'slug': slug, 'company': name, 'scope': scope, 'resumed': False,
              'started_at': datetime.now(timezone.utc).isoformat()}
    try:
        payload = tupu.collect(slug, scope, directory)
        coverage = payload.get('coverage') or {}
        record['status'] = coverage.get('status')
        record['complete'] = coverage.get('complete')
        record['pagination_exhausted'] = coverage.get('pagination_exhausted')
        record['detail_complete'] = coverage.get('detail_complete')
        record['collected_jobs'] = coverage.get('collected_jobs')
        record['expected_total'] = coverage.get('expected_total')
        record['requests'] = (coverage.get('request_budget') or {}).get('used')
        record['errors'] = list(coverage.get('errors') or [])
        record['note'] = coverage.get('note') or ''
        record['channels_used'] = list(coverage.get('channels_used') or [])
        validated = pipeline.validate_result(payload, name, scope)
        record['validation_ok'] = True
    except Exception as error:
        record['validation_ok'] = False
        record['status'] = record.get('status') or 'error'
        record['complete'] = False
        record['error'] = f'{type(error).__name__}: {error}'
        record.setdefault('errors', []).append(record['error'])
        record['traceback'] = traceback.format_exc()[-4000:]
        (directory / 'validate_error.json').write_text(
            json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
        record['elapsed_seconds'] = round(time.time() - started, 1)
        return record

    record['elapsed_seconds'] = round(time.time() - started, 1)
    record['finished_at'] = datetime.now(timezone.utc).isoformat()
    record['tenant_host'] = tupu.tenant_host(slug)
    record['list_url'] = tupu.list_url(slug, tupu.channel_for(slug, scope))
    marker.write_text(json.dumps({'slug': slug, 'company': name, 'scope': scope,
                                  'meta': record, 'validated': validated},
                                 ensure_ascii=False), encoding='utf-8')
    return record


def needs_retry(record, census_expected=None):
    """A unit worth a second pass: it failed, or it came back short of its own total.

    ``census_expected`` is the 2026-09-19 census count for that unit.  A blocked
    unit reports ``expected_total = 0`` (it read an empty list), so the census value
    is the only way to notice that a transient network error -- not an empty
    channel -- produced the zero.  茵梦达/intern was exactly that case.
    """
    if record.get('validation_ok') is False:
        return True
    if record.get('status') == 'error':
        return True
    collected = record.get('collected_jobs')
    expected = record.get('expected_total')
    if isinstance(expected, int) and isinstance(collected, int) and collected < expected:
        return True
    if not (collected or 0) and (census_expected or 0) > 0:
        return True
    if record.get('pagination_exhausted') is False and (collected or 0) > 0:
        return True
    return False


# --------------------------------------------------------------------------- #
# tables
# --------------------------------------------------------------------------- #
def demo_reason(name: str) -> str:
    text = str(name)
    if text.isdigit():
        return f'纯数字名称（{text}）—— 平台内部/压测租户，非企业'
    hits = [word for word in ('试用', '演示', '测试', '回归', '图谱') if word in text]
    if hits:
        return f'名称含「{"」「".join(hits)}」—— 演示/试用/回归租户，非真实企业'
    if text.lower().startswith('tupu'):
        return '图谱自有租户（tupu*）—— 平台演示环境，非企业客户'
    return '演示/试用/回归租户（站长 2026-09-19 判定口径）'


def build_tables(out: Path, inputs: Path, table_dir: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    scope, census, tenants = load_inputs(inputs)
    names = tenant_names(tenants)
    include = scope['include']
    exclude = scope['exclude']

    units = {}
    for directory in sorted((out / 'units').glob('*')):
        if not directory.is_dir():
            continue
        cached = read_validated(directory / 'validated.json')
        if cached is None:
            continue
        meta = cached.get('meta') or {}
        units[(meta.get('slug'), meta.get('scope'))] = cached

    jobs = []
    unit_rows = []
    for key in sorted(units, key=lambda k: (str(k[0]), str(k[1]))):
        slug, sc = key
        cached = units[key]
        meta = cached.get('meta') or {}
        record = cached.get('validated') or {}
        rows = record.get('jobs') or []
        company = meta.get('company') or (include.get(slug) or {}).get('name') or slug
        for row in rows:
            jobs.append((slug, company, sc, row))
        unit_rows.append({
            '公司名称': company,
            '租户slug': slug,
            '招聘类型': SCOPE_LABEL[sc],
            '状态': meta.get('status'),
            '完整': '是' if meta.get('complete') else '否',
            '采集条数': len(rows),
            '官方总数': meta.get('expected_total'),
            '请求数': meta.get('requests'),
            '耗时秒': meta.get('elapsed_seconds'),
            '分页到底': '是' if meta.get('pagination_exhausted') else '否',
            '详情完整': '是' if meta.get('detail_complete') else '否',
            '校验通过': '是' if meta.get('validation_ok') else '否',
            '取法': ' / '.join(meta.get('channels_used') or []),
            '错误数': len(meta.get('errors') or []),
            '首条错误': (meta.get('errors') or [''])[0][:400],
            '官方列表页': meta.get('list_url') or '',
        })

    # ---------------- sheet 1: every posting ---------------------------------
    job_header = ['公司名称', '租户slug', '招聘类型', '岗位名称', '工作城市', '职能类别',
                  '招聘人数', '发布时间', '截止时间', '学历要求', '岗位描述', '详情链接',
                  '投递链接', 'source_record_id', '岗位状态', '发布时间来源', '截止时间来源',
                  '岗位描述来源', '采集取法', '记录id']
    job_table = []
    for slug, company, sc, row in jobs:
        description = str(row.get('description_raw') or '')
        if len(description) > EXCEL_CELL_LIMIT:
            description = description[:EXCEL_CELL_LIMIT] + '…（超出 Excel 单元格上限，已截断）'
        job_table.append([
            company,
            slug,
            row.get('recruitment_type') or SCOPE_LABEL[sc],
            row.get('job_title') or '',
            row.get('city') or '',
            row.get('position_function_raw') or '',
            row.get('headcount_raw') or '',
            row.get('published_at') or '',
            row.get('deadline_raw') or '',
            row.get('education_raw') or '',
            description,
            row.get('detail_url') or '',
            row.get('application_url') or row.get('detail_url') or '',
            row.get('source_record_id') or '',
            row.get('source_status_raw') or '',
            row.get('published_at_source') or '',
            row.get('deadline_source') or '',
            row.get('description_source') or '',
            row.get('channel_used') or '',
            row.get('id') or '',
        ])

    # ---------------- sheet 2: tenant summary --------------------------------
    per_tenant = {}
    for key, cached in units.items():
        slug, sc = key
        meta = cached.get('meta') or {}
        rows = (cached.get('validated') or {}).get('jobs') or []
        entry = per_tenant.setdefault(slug, {
            '公司名称': meta.get('company') or (include.get(slug) or {}).get('name') or slug,
            '租户host': meta.get('tenant_host') or '',
            'requests': 0, 'jobs': 0, 'published': 0, 'deadline': 0,
            'scopes': {}, 'errors': 0,
        })
        entry['requests'] += int(meta.get('requests') or 0)
        entry['jobs'] += len(rows)
        entry['published'] += sum(1 for row in rows if row.get('published_at'))
        entry['deadline'] += sum(1 for row in rows if row.get('deadline_raw'))
        entry['errors'] += len(meta.get('errors') or [])
        entry['scopes'][sc] = meta
    tenant_header = ['公司名称', '租户slug', '租户host',
                     '校招条数', '校招状态', '实习条数', '实习状态', '社招条数', '社招状态',
                     '合计条数', '总请求数', '发布时间可用率', '截止时间可用率', '错误条数']
    tenant_table = []
    for slug in sorted(per_tenant, key=lambda s: (-per_tenant[s]['jobs'], s)):
        entry = per_tenant[slug]

        def status_of(sc):
            return (entry['scopes'].get(sc) or {}).get('status')

        def count_of(sc):
            return len(((units.get((slug, sc)) or {}).get('validated') or {}).get('jobs') or [])

        total = entry['jobs']
        tenant_table.append([
            entry['公司名称'], slug, entry['租户host'],
            count_of('campus'), status_of('campus'),
            count_of('intern'), status_of('intern'),
            count_of('social'), status_of('social'),
            total, entry['requests'],
            f"{entry['published']}/{total}" if total else '0/0',
            f"{entry['deadline']}/{total}" if total else '0/0',
            entry['errors'],
        ])

    # ---------------- sheet 3: unit detail -----------------------------------
    unit_header = ['公司名称', '租户slug', '招聘类型', '状态', '完整', '采集条数', '官方总数',
                   '请求数', '耗时秒', '分页到底', '详情完整', '校验通过', '取法', '错误数',
                   '首条错误', '官方列表页']

    # ---------------- sheet 4: exclusions ------------------------------------
    exclude_rows = []
    for slug in sorted(exclude, key=lambda s: -int((exclude[s] or {}).get('postings') or 0)):
        entry = exclude[slug] or {}
        name = entry.get('name') or slug
        exclude_rows.append(['演示/试用租户', name, slug,
                             entry.get('postings'), demo_reason(name),
                             '2026-09-19 全站普查（三频道合计）'])
    zero_targets = sorted(set(census) - set(include) - set(exclude))
    for slug in zero_targets:
        note = '2026-09-19 全站普查：三频道均 0 条'
        if slug == 'bbac':
            note = ('2026-09-19 普查返回 194 字节无法归类；本次重试 3 频道官方列表页 + '
                    'nextPageList 均返回空分页器（195 字节），确认 0 条')
        name = names.get(slug) or (scope.get('retry') or {}).get(slug) or slug
        exclude_rows.append(['0 职位租户', name, slug, 0,
                             '官方公开频道 0 条职位（非演示租户，仅当前无发布）', note])
    exclude_header = ['类别', '名称', '租户slug', '三频道条数', '排除理由', '证据/备注']

    # ---------------- sheet 5: accounting ------------------------------------
    total_requests = sum(int((cached.get('meta') or {}).get('requests') or 0)
                         for cached in units.values())
    total_jobs = sum(len((cached.get('validated') or {}).get('jobs') or [])
                     for cached in units.values())
    error_rows = []
    for key in sorted(units, key=lambda k: (str(k[0]), str(k[1]))):
        slug, sc = key
        meta = units[key].get('meta') or {}
        for message in meta.get('errors') or []:
            error_rows.append([meta.get('company') or slug, slug, SCOPE_LABEL[sc], message])
    started = [(cached.get('meta') or {}).get('started_at') for cached in units.values()]
    finished = [(cached.get('meta') or {}).get('finished_at') for cached in units.values()]
    elapsed = sum(float((cached.get('meta') or {}).get('elapsed_seconds') or 0)
                  for cached in units.values())
    accounting = [
        ['采集目标数（租户 × 频道）', len(units)],
        ['采集目标租户数', len({slug for slug, _ in units})],
        ['岗位总条数', total_jobs],
        ['公开只读请求总数', total_requests],
        ['单元耗时合计（秒）', round(elapsed, 1)],
        ['单元耗时合计（小时）', round(elapsed / 3600, 2)],
        ['单请求间隔（秒，QIUZHAO_PLATFORM_REQUEST_INTERVAL）', DEFAULT_INTERVAL],
        ['同 host 并发', MAX_WORKERS],
        ['聚合速率上限（req/s）', round(MAX_WORKERS / DEFAULT_INTERVAL, 2)],
        ['最早单元开始时间（UTC）', min([x for x in started if x] or [''])],
        ['最晚单元完成时间（UTC）', max([x for x in finished if x] or [''])],
        ['平均每单元请求数', round(total_requests / len(units), 2) if units else 0],
        ['错误消息条数', len(error_rows)],
        ['失败/未通过校验的单元数',
         sum(1 for cached in units.values() if not (cached.get('meta') or {}).get('validation_ok'))],
    ]

    # ---------------- write workbook -----------------------------------------
    wb = Workbook()
    head_font = Font(bold=True, color='FFFFFF')
    head_fill = PatternFill('solid', fgColor='305496')
    wrap = Alignment(vertical='top', wrap_text=True)
    top = Alignment(vertical='top')

    def sheet(name, header, rows, wrap_columns=()):
        # A sheet may be handed dicts (unit_rows) or lists; normalise to header order.
        rows = [[row.get(column) for column in header] if isinstance(row, dict) else list(row)
                for row in rows]
        ws = wb.create_sheet(name) if wb.sheetnames != ['Sheet'] else wb.active
        ws.title = name
        ws.append(header)
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
        for row in rows:
            ws.append(row)
        for index, column in enumerate(header, start=1):
            letter = get_column_letter(index)
            sample = [len(str(row[index - 1] or '')) for row in rows[:200]]
            width = max(10, min(48, max([len(str(column)) * 2] + sample + [10]) + 2))
            ws.column_dimensions[letter].width = width
        for index in wrap_columns:
            for cell in ws[get_column_letter(index)][1:]:
                cell.alignment = wrap
        ws.freeze_panes = 'A2'
        return ws

    sheet('全部岗位', job_header, job_table, wrap_columns=(11,))
    sheet('租户汇总', tenant_header, tenant_table)
    sheet('单元明细', unit_header, unit_rows)
    sheet('排除名单', exclude_header, exclude_rows)
    account_rows = [[key, value] for key, value in accounting]
    account_rows.append(['', ''])
    account_rows.append(['--- 错误清单（每单元逐条） ---', ''])
    account_rows.extend(error_rows)
    sheet('采集记账', ['项目', '值'], account_rows, wrap_columns=(2,))

    table_dir.mkdir(parents=True, exist_ok=True)
    xlsx = table_dir / f'{TABLE_STEM}.xlsx'
    wb.save(xlsx)

    csv_path = table_dir / f'{TABLE_STEM}.csv'
    with csv_path.open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(job_header)
        writer.writerows(job_table)

    summary = {'units': len(units), 'jobs': total_jobs, 'requests': total_requests,
               'errors': len(error_rows), 'xlsx': str(xlsx), 'csv': str(csv_path),
               'sheets': wb.sheetnames}
    (out / 'table-summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2),
                                            encoding='utf-8')
    return summary


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #
def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out', type=Path, default=DEFAULT_OUT)
    parser.add_argument('--inputs', type=Path, default=INPUTS)
    parser.add_argument('--table-dir', type=Path, default=TABLE_DIR)
    parser.add_argument('--workers', type=int, default=MAX_WORKERS)
    parser.add_argument('--interval', type=float, default=DEFAULT_INTERVAL)
    parser.add_argument('--passes', type=int, default=2,
                        help='total collection passes; pass 2 only re-runs failed/short units')
    parser.add_argument('--only', default='', help='comma separated tenant slugs')
    parser.add_argument('--scopes', default=','.join(SCOPES))
    parser.add_argument('--force', action='store_true', help='ignore existing validated.json')
    parser.add_argument('--build-table', action='store_true',
                        help='skip collection, only build tables')
    args = parser.parse_args()

    if args.interval < DEFAULT_INTERVAL:
        parser.error(f'--interval must be >= {DEFAULT_INTERVAL} (hard constraint)')
    workers = max(1, min(MAX_WORKERS, args.workers))

    scope_data, _census, _tenants = load_inputs(args.inputs)
    include = scope_data['include']
    slugs = sorted(include)
    if args.only:
        wanted = {item.strip() for item in args.only.split(',') if item.strip()}
        slugs = [slug for slug in slugs if slug in wanted]
    scopes = [item.strip() for item in args.scopes.split(',') if item.strip()]
    for sc in scopes:
        if sc not in SCOPES:
            parser.error(f'unknown scope {sc}')

    args.out.mkdir(parents=True, exist_ok=True)
    (args.out / 'logs').mkdir(parents=True, exist_ok=True)

    os.environ['QIUZHAO_PLATFORM_REQUEST_INTERVAL'] = str(args.interval)
    os.environ.pop('QIUZHAO_TUPU360_HEADLESS', None)  # documented ladder only

    ledger_path = args.out / 'logs' / 'unit-ledger.jsonl'
    ledger_lock = threading.Lock()

    census_by_unit = {(slug, sc): int((include[slug].get(CENSUS_CHANNEL[sc]) or {}).get('n') or 0)
                      for slug in slugs for sc in scopes}

    def run_pass(pass_no, plan):
        log(f'pass {pass_no}: {len(plan)} units, workers={workers}, interval={args.interval}s')
        done = 0
        retry_plan = []
        with ThreadPoolExecutor(max_workers=workers) as pool:
            futures = {pool.submit(collect_unit, slug, include[slug]['name'], sc, args.out,
                                   args.force and pass_no == 1): (slug, sc)
                       for slug, sc in plan}
            for future in as_completed(futures):
                slug, sc = futures[future]
                try:
                    record = future.result()
                except Exception as error:  # defensive: collect_unit already catches
                    record = {'slug': slug, 'scope': sc, 'validation_ok': False,
                              'status': 'error',
                              'errors': [f'{type(error).__name__}: {error}']}
                done += 1
                with ledger_lock:
                    with ledger_path.open('a', encoding='utf-8') as handle:
                        handle.write(json.dumps(record, ensure_ascii=False) + '\n')
                log(f'  [{done}/{len(plan)}] {slug}/{sc} {record.get("status")} '
                    f'jobs={record.get("collected_jobs")} req={record.get("requests")} '
                    f'{record.get("elapsed_seconds")}s '
                    f'{"OK" if record.get("validation_ok") else "VALIDATE-FAIL"}')
                if pass_no < args.passes and needs_retry(record, census_by_unit.get((slug, sc))):
                    retry_plan.append((slug, sc))
        return retry_plan

    def expected_jobs(slug, sc):
        return int((include[slug].get(CENSUS_CHANNEL[sc]) or {}).get('n') or 0)

    if not args.build_table:
        # Largest units first: a single 828-posting unit is one worker's whole slot,
        # so starting it last would leave a 25-minute tail on an otherwise idle pool.
        plan = sorted(((slug, sc) for slug in slugs for sc in scopes),
                      key=lambda unit: (-expected_jobs(*unit), unit[0], unit[1]))
        for pass_no in range(1, args.passes + 1):
            if not plan:
                break
            plan = run_pass(pass_no, plan)
            if pass_no < args.passes:
                log(f'pass {pass_no} leftovers to re-run: {len(plan)}')
                for slug, sc in plan:
                    marker = unit_dir(args.out, slug, sc) / 'validated.json'
                    if marker.exists():
                        marker.unlink()

    summary = build_tables(args.out, args.inputs, args.table_dir)
    log('summary: ' + json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
